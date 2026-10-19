from pathlib import Path

import pytest
from docx import Document

from backend.app.rag.parsing.documents import (
    UnsupportedDocumentType,
    parse_document_text,
)


def _write_text_pdf(path: Path, text: str) -> None:
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        (
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            b"/Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>"
        ),
    ]
    stream = ("BT /F1 12 Tf 72 720 Td ({0}) Tj ET".format(text)).encode("ascii")
    objects.append(
        b"<< /Length " + str(len(stream)).encode("ascii") + b" >>\nstream\n"
        + stream
        + b"\nendstream"
    )
    objects.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")

    payload = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for number, body in enumerate(objects, start=1):
        offsets.append(len(payload))
        payload.extend("{0} 0 obj\n".format(number).encode("ascii"))
        payload.extend(body)
        payload.extend(b"\nendobj\n")

    xref_offset = len(payload)
    payload.extend("xref\n0 {0}\n".format(len(objects) + 1).encode("ascii"))
    payload.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        payload.extend("{0:010d} 00000 n \n".format(offset).encode("ascii"))
    payload.extend(
        (
            "trailer\n<< /Size {0} /Root 1 0 R >>\nstartxref\n{1}\n%%EOF\n".format(
                len(objects) + 1, xref_offset
            )
        ).encode("ascii")
    )
    path.write_bytes(bytes(payload))


def test_pdf_parser_preserves_page_sections(tmp_path):
    path = tmp_path / "manual.pdf"
    _write_text_pdf(path, "PDF knowledge")

    content = parse_document_text(path)

    assert "# Page 1" in content
    assert "PDF knowledge" in content


def test_docx_parser_preserves_heading_structure(tmp_path):
    path = tmp_path / "manual.docx"
    document = Document()
    document.add_heading("Deployment", level=1)
    document.add_paragraph("Use the local Docker stack.")
    document.save(path)

    content = parse_document_text(path)

    assert "# Deployment" in content
    assert "Use the local Docker stack." in content


def test_parser_rejects_unsupported_documents(tmp_path):
    path = tmp_path / "archive.zip"
    path.write_bytes(b"not a document")

    with pytest.raises(UnsupportedDocumentType):
        parse_document_text(path)
