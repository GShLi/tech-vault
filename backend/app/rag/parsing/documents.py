"""Convert supported source documents into structure-aware Markdown text."""

from __future__ import annotations

import re
from pathlib import Path
from typing import Union


class UnsupportedDocumentType(ValueError):
    """Raised when no parser is registered for a source document."""


_HEADING_STYLE = re.compile(r"^(?:Heading|标题)\s*([1-6])$", re.IGNORECASE)


def _read_plain_text(path: Path) -> str:
    return path.read_text(encoding="utf-8-sig")


def _read_pdf(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    pages = []
    for page_number, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        if text:
            pages.append("# Page {0}\n\n{1}".format(page_number, text))
    return "\n\n".join(pages)


def _docx_paragraph_markdown(paragraph) -> str:
    text = paragraph.text.strip()
    if not text:
        return ""
    style_name = paragraph.style.name if paragraph.style is not None else ""
    match = _HEADING_STYLE.match(style_name)
    if match is not None:
        return "{0} {1}".format("#" * int(match.group(1)), text)
    if style_name.casefold() == "title":
        return "# " + text
    return text


def _docx_table_markdown(table) -> str:
    rows = []
    for row in table.rows:
        cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
        if any(cells):
            rows.append("| " + " | ".join(cells) + " |")
    return "\n".join(rows)


def _read_docx(path: Path) -> str:
    from docx import Document

    document = Document(str(path))
    blocks = []
    for paragraph in document.paragraphs:
        rendered = _docx_paragraph_markdown(paragraph)
        if rendered:
            blocks.append(rendered)
    for table in document.tables:
        rendered = _docx_table_markdown(table)
        if rendered:
            blocks.append(rendered)
    return "\n\n".join(blocks)


def parse_document_text(path: Union[str, Path]) -> str:
    """Return Markdown-compatible text for a supported source document."""

    source = Path(path)
    suffix = source.suffix.lower()
    if suffix in {".md", ".txt"}:
        return _read_plain_text(source)
    if suffix == ".pdf":
        return _read_pdf(source)
    if suffix == ".docx":
        return _read_docx(source)
    raise UnsupportedDocumentType("Unsupported document type: " + suffix)
