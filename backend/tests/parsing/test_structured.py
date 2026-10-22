from pathlib import Path

import pymupdf
import pytest
from docx import Document

from backend.app.rag.parsing.router import parse_document
from backend.app.rag.parsing.contracts import NodeType, ParseError


def test_docx_keeps_order(tmp_path):
    path = tmp_path / "order.docx"
    doc = Document()
    doc.add_heading("Guide", 1)
    doc.add_paragraph("Before")
    doc.add_table(rows=1, cols=1).cell(0, 0).text = "Table"
    doc.add_paragraph("After")
    doc.save(path)
    result = parse_document(path)
    assert [node.node_type for node in result.nodes] == [NodeType.HEADING, NodeType.PARAGRAPH, NodeType.TABLE, NodeType.PARAGRAPH]


def test_txt_has_no_markdown_structure_and_fence_heading_is_code(tmp_path):
    path = tmp_path / "plain.txt"
    path.write_text("# literal\n\n```literal\ntext", encoding="utf-8")
    assert all(node.node_type == NodeType.PARAGRAPH for node in parse_document(path).nodes)
    path = tmp_path / "code.md"
    path.write_text("# Guide\n\n````python\n# comment\n```\n````\nNext paragraph", encoding="utf-8")
    nodes = parse_document(path).nodes
    assert [node.node_type for node in nodes] == [NodeType.HEADING, NodeType.CODE, NodeType.PARAGRAPH]


def test_native_pdf_and_scanned_rejection(tmp_path):
    path = tmp_path / "native.pdf"
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((60, 90), "Spring Boot", fontsize=20)
    page.insert_text((60, 130), "Install Java 21 before starting the service. " * 2)
    doc.save(path)
    doc.close()
    result = parse_document(path)
    assert result.pdf_kind == "native"
    assert any(node.page_no == 1 and node.bbox for node in result.nodes)
    scan = tmp_path / "scan.pdf"
    doc = pymupdf.open()
    page = doc.new_page()
    pix = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 20, 20), False)
    pix.clear_with(240)
    page.insert_image(page.rect, pixmap=pix)
    doc.save(scan)
    doc.close()
    with pytest.raises(ParseError, match="PDF_SCANNED_UNSUPPORTED"):
        parse_document(scan)


def test_mime_mismatch_is_rejected(tmp_path):
    path = tmp_path / "a.md"
    path.write_text("text")
    with pytest.raises(ParseError, match="DOCUMENT_TYPE_MISMATCH"):
        parse_document(path, mime_type="application/pdf")
