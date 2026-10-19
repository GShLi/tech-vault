from dataclasses import replace
import hashlib
from io import BytesIO

from docx import Document as DocxDocument
from sqlalchemy import select

from backend.app.db.models import ChunkRecord, Document, DocumentVersion
from backend.tests.api.conftest import admin_headers_for


def _create_team_and_kb(client, headers):
    team = client.post(
        "/api/v1/teams/create",
        json={"name": "DocTest"},
        headers=headers,
    ).json()["data"]
    kb = client.post(
        "/api/v1/kbs/create",
        json={"team_id": team["id"], "name": "Test KB"},
        headers=headers,
    ).json()["data"]
    return team, kb


def test_upload_returns_202_and_lists_document(api_client):
    headers = admin_headers_for(api_client)
    team, kb = _create_team_and_kb(api_client, headers)

    upload = api_client.post(
        "/api/v1/docs/upload",
        data={
            "team_id": team["id"],
            "kb_id": kb["id"],
            "visibility": "team_shared",
            "data_classification": "internal",
        },
        files={"file": ("README.md", b"# Hello\n\nThis is a test document.")},
        headers=headers,
    )

    assert upload.status_code == 202
    assert upload.json()["data"]["document"]["file_name"] == "README.md"
    assert upload.json()["data"]["document"]["doc_status"] == "active"
    assert upload.json()["data"]["document_version_id"]
    assert upload.json()["data"]["ingestion_run_id"]

    doc_list = api_client.get(
        "/api/v1/docs",
        params={"team_id": team["id"], "kb_id": kb["id"]},
        headers=headers,
    )
    assert doc_list.status_code == 200
    assert len(doc_list.json()["data"]) == 1

    detail = api_client.get(
        "/api/v1/docs/detail",
        params={"team_id": team["id"], "document_id": upload.json()["data"]["document"]["id"]},
        headers=headers,
    )
    assert detail.status_code == 200
    assert detail.json()["data"]["file_name"] == "README.md"

    ingestion_list = api_client.get(
        "/api/v1/ingestions",
        params={"team_id": team["id"]},
        headers=headers,
    )
    assert ingestion_list.status_code == 200
    assert len(ingestion_list.json()["data"]) == 1

    ingestion_detail = api_client.get(
        "/api/v1/ingestions/detail",
        params={"team_id": team["id"], "run_id": upload.json()["data"]["ingestion_run_id"]},
        headers=headers,
    )
    assert ingestion_detail.status_code == 200
    assert ingestion_detail.json()["data"]["status"] == "completed"


def test_upload_accepts_top_level_metadata_and_fixes_owner(api_client):
    admin_headers = admin_headers_for(api_client)
    team, kb = _create_team_and_kb(api_client, admin_headers)
    me = api_client.get("/api/v1/auth/me", headers=admin_headers).json()["data"]

    upload = api_client.post(
        "/api/v1/docs/upload",
        data={
            "team_id": team["id"],
            "kb_id": kb["id"],
            "product": "TechVault",
            "version_values": ["1.0", "2.0"],
            "visibility": "team_shared",
            "data_classification": "internal",
            "metadata": '{"owner_user_id":"forged-user","product":"Ignored"}',
        },
        files={"file": ("metadata.md", b"# Metadata\n\nBody\n", "text/markdown")},
        headers=admin_headers,
    )

    assert upload.status_code == 202, upload.text
    document = upload.json()["data"]["document"]
    detail = api_client.get(
        "/api/v1/docs/detail",
        params={"team_id": team["id"], "document_id": document["id"]},
        headers=admin_headers,
    )
    assert detail.status_code == 200, detail.text
    body = detail.json()["data"]
    assert body["product"] == "TechVault"
    assert body["version_values"] == ["1.0", "2.0"]
    assert body["owner_user_id"] == me["user"]["id"]


def test_document_update_rejects_invalid_version_values(api_client):
    headers = admin_headers_for(api_client)
    team, kb = _create_team_and_kb(api_client, headers)
    upload = api_client.post(
        "/api/v1/docs/upload",
        data={"team_id": team["id"], "kb_id": kb["id"]},
        files={"file": ("versions.md", b"# Versions\n\nBody\n")},
        headers=headers,
    )
    assert upload.status_code == 202, upload.text
    document = upload.json()["data"]["document"]

    invalid = api_client.post(
        "/api/v1/docs/update",
        json={
            "team_id": team["id"],
            "document_id": document["id"],
            "version_values": "3.2",
        },
        headers={**headers, "If-Match": str(document["version"])},
    )
    assert invalid.status_code == 400
    assert invalid.json()["error"]["code"] == "INVALID_ARGUMENT"


def test_document_update_rejects_invalid_visibility_type(api_client):
    headers = admin_headers_for(api_client)
    team, kb = _create_team_and_kb(api_client, headers)
    upload = api_client.post(
        "/api/v1/docs/upload",
        data={"team_id": team["id"], "kb_id": kb["id"]},
        files={"file": ("visibility.md", b"# Visibility\n\nBody\n")},
        headers=headers,
    )
    assert upload.status_code == 202, upload.text
    document = upload.json()["data"]["document"]

    invalid = api_client.post(
        "/api/v1/docs/update",
        json={
            "team_id": team["id"],
            "document_id": document["id"],
            "visibility": [],
        },
        headers={**headers, "If-Match": str(document["version"])},
    )

    assert invalid.status_code == 400
    assert invalid.json()["error"]["code"] == "INVALID_ARGUMENT"



def test_knowledge_base_document_count_excludes_deleted_documents(api_client):
    headers = admin_headers_for(api_client)
    team, kb = _create_team_and_kb(api_client, headers)
    upload = api_client.post(
        "/api/v1/docs/upload",
        data={"team_id": team["id"], "kb_id": kb["id"]},
        files={"file": ("counted.md", b"# Counted\n\nCount me once.")},
        headers=headers,
    )
    assert upload.status_code == 202, upload.text
    document = upload.json()["data"]["document"]

    listing = api_client.get(
        "/api/v1/kbs",
        params={"team_id": team["id"]},
        headers=headers,
    )
    detail = api_client.get(
        "/api/v1/kbs/detail",
        params={"team_id": team["id"], "kb_id": kb["id"]},
        headers=headers,
    )
    assert listing.status_code == 200
    assert detail.status_code == 200
    assert listing.json()["data"][0]["document_count"] == 1
    assert detail.json()["data"]["document_count"] == 1

    deleted = api_client.post(
        "/api/v1/docs/delete",
        json={"team_id": team["id"], "document_id": document["id"]},
        headers=headers,
    )
    assert deleted.status_code == 202, deleted.text
    listing = api_client.get(
        "/api/v1/kbs",
        params={"team_id": team["id"]},
        headers=headers,
    )
    detail = api_client.get(
        "/api/v1/kbs/detail",
        params={"team_id": team["id"], "kb_id": kb["id"]},
        headers=headers,
    )
    assert listing.json()["data"][0]["document_count"] == 0
    assert detail.json()["data"]["document_count"] == 0


def test_upload_version_fingerprint_is_sha256_of_uploaded_bytes(api_client, api_app):
    headers = admin_headers_for(api_client)
    team, kb = _create_team_and_kb(api_client, headers)
    original = b"# Fingerprint\n\nExact uploaded content."

    upload = api_client.post(
        "/api/v1/docs/upload",
        data={"team_id": team["id"], "kb_id": kb["id"]},
        files={"file": ("fingerprint.md", original)},
        headers=headers,
    )

    assert upload.status_code == 202, upload.text
    document_id = upload.json()["data"]["document"]["id"]
    with api_app.state.session_factory() as session:
        version = session.scalar(
            select(DocumentVersion).where(DocumentVersion.document_id == document_id)
        )
    assert version is not None
    assert version.content_fingerprint == hashlib.sha256(original).hexdigest()


def test_update_document_rejects_path_in_file_name(api_client):
    headers = admin_headers_for(api_client)
    team, kb = _create_team_and_kb(api_client, headers)
    upload = api_client.post(
        "/api/v1/docs/upload",
        data={"team_id": team["id"], "kb_id": kb["id"]},
        files={"file": ("safe.md", b"# Safe\n\nOriginal name.")},
        headers=headers,
    )
    assert upload.status_code == 202, upload.text
    document = upload.json()["data"]["document"]

    response = api_client.post(
        "/api/v1/docs/update",
        json={
            "team_id": team["id"],
            "document_id": document["id"],
            "file_name": "../escape.md",
        },
        headers={**headers, "If-Match": str(document["version"])},
    )

    assert response.status_code == 400, response.text
    assert response.json()["error"]["code"] == "INVALID_ARGUMENT"
    detail = api_client.get(
        "/api/v1/docs/detail",
        params={"team_id": team["id"], "document_id": document["id"]},
        headers=headers,
    )
    assert detail.json()["data"]["file_name"] == "safe.md"


def test_download_cache_path_does_not_use_untrusted_file_name(
    api_client, api_app, monkeypatch, tmp_path
):
    headers = admin_headers_for(api_client)
    team, kb = _create_team_and_kb(api_client, headers)
    upload = api_client.post(
        "/api/v1/docs/upload",
        data={"team_id": team["id"], "kb_id": kb["id"]},
        files={"file": ("safe.md", b"# Safe\n\nDownload content.")},
        headers=headers,
    )
    assert upload.status_code == 202, upload.text
    document_id = upload.json()["data"]["document"]["id"]
    upload_dir = tmp_path / "uploads"
    api_app.state.settings = replace(api_app.state.settings, upload_dir=str(upload_dir))

    with api_app.state.session_factory() as session:
        document = session.get(Document, document_id)
        document.source_uri = str(tmp_path / "missing-source.md")
        document.object_uri = "s3://techvault/documents/cache-test"
        document.file_name = "../../unsafe.md"
        session.commit()

    captured = {}

    def fake_download_s3_object(source_uri, destination, settings):
        captured["source_uri"] = source_uri
        captured["destination"] = destination
        destination.parent.mkdir(parents=True, exist_ok=True)
        destination.write_bytes(b"# Safe\n\nDownload content.")
        return destination

    monkeypatch.setattr(
        "backend.app.api.v1.documents.download_s3_object",
        fake_download_s3_object,
    )

    response = api_client.get(
        "/api/v1/docs/download",
        params={"team_id": team["id"], "document_id": document_id},
        headers=headers,
    )

    assert response.status_code == 200, response.text
    destination = captured["destination"]
    cache_root = (upload_dir / "cache").resolve()
    assert destination.resolve().is_relative_to(cache_root)
    assert destination.name == document_id + ".download"
    assert "unsafe.md" not in destination.name


def test_upload_rejects_invalid_visibility(api_client):
    headers = admin_headers_for(api_client)
    team, kb = _create_team_and_kb(api_client, headers)

    response = api_client.post(
        "/api/v1/docs/upload",
        data={
            "team_id": team["id"],
            "kb_id": kb["id"],
            "visibility": "public",
        },
        files={"file": ("invalid.md", b"# Invalid\n\nVisibility.")},
        headers=headers,
    )

    assert response.status_code == 400
    assert response.json()["error"]["code"] == "INVALID_ARGUMENT"


def test_uploaded_chunks_keep_team_and_knowledge_base_scope(api_client, api_app):
    headers = admin_headers_for(api_client)
    team, kb = _create_team_and_kb(api_client, headers)

    upload = api_client.post(
        "/api/v1/docs/upload",
        data={"team_id": team["id"], "kb_id": kb["id"]},
        files={"file": ("scope.md", b"# Scope\n\nScoped content.")},
        headers=headers,
    )
    assert upload.status_code == 202

    with api_app.state.session_factory() as session:
        chunks = session.scalars(select(ChunkRecord)).all()
        assert chunks
        assert all(chunk.team_id == team["id"] for chunk in chunks)
        assert all(chunk.knowledge_base_id == kb["id"] for chunk in chunks)


def test_upload_processes_docx_document(api_client):
    headers = admin_headers_for(api_client)
    team, kb = _create_team_and_kb(api_client, headers)
    payload = BytesIO()
    docx = DocxDocument()
    docx.add_heading("Runbook", level=1)
    docx.add_paragraph("Restart the retrieval worker.")
    docx.save(payload)

    upload = api_client.post(
        "/api/v1/docs/upload",
        data={"team_id": team["id"], "kb_id": kb["id"]},
        files={
            "file": (
                "runbook.docx",
                payload.getvalue(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
        headers=headers,
    )

    assert upload.status_code == 202
    assert upload.json()["data"]["document"]["doc_status"] == "active"
    detail = api_client.get(
        "/api/v1/ingestions/detail",
        params={
            "team_id": team["id"],
            "run_id": upload.json()["data"]["ingestion_run_id"],
        },
        headers=headers,
    )
    assert detail.json()["data"]["status"] == "completed"


def test_document_endpoints_reject_another_team_scope(api_client):
    owner_headers = admin_headers_for(api_client)
    other_headers = owner_headers

    owner_team, owner_kb = _create_team_and_kb(api_client, owner_headers)
    other_team = api_client.post(
        "/api/v1/teams/create",
        json={"name": "OtherTeam"},
        headers=other_headers,
    ).json()["data"]

    upload = api_client.post(
        "/api/v1/docs/upload",
        data={"team_id": owner_team["id"], "kb_id": owner_kb["id"]},
        files={"file": ("private.md", b"# Private\n\nOwner content.")},
        headers=owner_headers,
    )
    assert upload.status_code == 202
    document_id = upload.json()["data"]["document"]["id"]

    cross_team = api_client.get(
        "/api/v1/docs/detail",
        params={"team_id": other_team["id"], "document_id": document_id},
        headers=other_headers,
    )
    assert cross_team.status_code == 404

    bad_upload = api_client.post(
        "/api/v1/docs/upload",
        data={"team_id": other_team["id"], "kb_id": owner_kb["id"]},
        files={"file": ("bad.md", b"# Bad\n\nWrong scope.")},
        headers=other_headers,
    )
    assert bad_upload.status_code == 404
