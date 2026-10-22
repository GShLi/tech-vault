"""Download official Java sources and record their provenance for live QA."""
import argparse
import hashlib
import json
from datetime import datetime, timezone
from html.parser import HTMLParser
from pathlib import Path

import httpx


class ArticleText(HTMLParser):
    def __init__(self):
        super().__init__()
        self.parts = []
        self.skipped = 0

    def handle_starttag(self, tag, attrs):
        if tag in {'script', 'style'}:
            self.skipped += 1
        if self.skipped:
            return
        if tag in {'h1', 'h2', 'h3', 'h4', 'p', 'div', 'li', 'tr'}:
            self.parts.append('\n\n')
        if tag in {'h1', 'h2', 'h3', 'h4'}:
            self.parts.append('#' * int(tag[1]) + ' ')
        if tag == 'pre':
            self.parts.append('\n\n```java\n')
        if tag == 'br':
            self.parts.append('\n')
        if tag == 'td':
            self.parts.append(' | ')

    def handle_endtag(self, tag):
        if tag in {'script', 'style'}:
            self.skipped = max(0, self.skipped - 1)
        if tag == 'pre':
            self.parts.append('\n```\n\n')

    def handle_data(self, data):
        if not self.skipped:
            self.parts.append(data)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--output', type=Path, default=Path('data/java-acceptance/corpus'))
    args = parser.parse_args()
    args.output.mkdir(parents=True, exist_ok=True)
    sources = [
        ('java-21-language-updates.pdf', 'https://docs.oracle.com/en/java/javase/21/language/java-se-language-updates.pdf'),
        ('jep-444-virtual-threads.md', 'https://openjdk.org/jeps/444'),
        ('jep-395-records.md', 'https://openjdk.org/jeps/395'),
        ('jep-409-sealed-classes.txt', 'https://openjdk.org/jeps/409'),
    ]
    manifest = []
    with httpx.Client(follow_redirects=True, timeout=90, trust_env=False) as client:
        for filename, url in sources:
            response = client.get(url)
            response.raise_for_status()
            original = response.content
            if filename.endswith('.pdf'):
                content = original
                conversion = 'original PDF'
            else:
                html = ArticleText()
                html.feed(response.text)
                content = (f'Source: {url}\n\n' + ''.join(html.parts)).encode()
                conversion = 'HTML converted to text/Markdown; code blocks retained'
                (args.output / (filename + '.source.html')).write_bytes(original)
            (args.output / filename).write_bytes(content)
            manifest.append({
                'file': filename, 'source': url, 'final_url': str(response.url),
                'downloaded_at': datetime.now(timezone.utc).isoformat(),
                'bytes': len(content), 'sha256': hashlib.sha256(content).hexdigest(),
                'source_sha256': hashlib.sha256(original).hexdigest(), 'conversion': conversion,
            })
            print(filename, len(content), 'bytes', flush=True)
    # A derived DOCX exercises ordered paragraphs, code and a table. It is clearly
    # identified as a test rendition, not an original OpenJDK publication.
    from docx import Document
    doc = Document()
    doc.add_heading('Java Records — OpenJDK JEP 395 test rendition', 0)
    doc.add_paragraph('Source: https://openjdk.org/jeps/395')
    doc.add_paragraph('Derived test rendition; consult the linked original for authoritative wording.')
    text = (args.output / 'jep-395-records.md').read_text()
    for paragraph in text.split('\n\n'):
        if paragraph.strip():
            doc.add_paragraph(paragraph.strip())
    doc.add_heading('QA fixture: record example', 1)
    doc.add_paragraph('record Point(int x, int y) {}')
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = 'Component'
    table.rows[0].cells[1].text = 'Accessor'
    for component, accessor in [('x', 'x()'), ('y', 'y()')]:
        cells = table.add_row().cells
        cells[0].text, cells[1].text = component, accessor
    path = args.output / 'java-records-derived-test.docx'
    doc.save(path)
    manifest.append({'file': path.name, 'source': 'https://openjdk.org/jeps/395',
                     'conversion': 'Derived DOCX with separately labelled QA table',
                     'bytes': path.stat().st_size, 'sha256': hashlib.sha256(path.read_bytes()).hexdigest()})
    (args.output / 'manifest.json').write_text(json.dumps(manifest, ensure_ascii=False, indent=2))


if __name__ == '__main__':
    main()
