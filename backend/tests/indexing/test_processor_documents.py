import json
from urllib.parse import parse_qs, urlparse

from docx import Document as DocxDocument

import backend.app.rag.ingestion.processor as processor_module
from backend.app.core.config import Settings
from backend.app.db.models import (
    ChunkRecord,
    Document,
    DocumentVersion,
    IngestionRun,
    KnowledgeBase,
    Team,
    TeamMember,
    User,
)
from backend.app.db.session import build_engine, build_session_factory, create_schema
from backend.app.rag.ingestion.processor import process_document


def _ingestion_context(tmp_path, file_name):
    engine = build_engine("sqlite+pysqlite:///:memory:")
    create_schema(engine)
    session_factory = build_session_factory(engine)
    settings = Settings(
        environment="test",
        jwt_secret="test-secret-at-least-32-bytes-long",
        enable_indexing=False,
    )
    session = session_factory()

    user = User(email="owner@example.com", display_name="Owner", password_hash="hash")
    session.add(user)
    session.flush()
    team = Team(name="Parsing Team", created_by=user.id)
    session.add(team)
    session.flush()
    session.add(
        TeamMember(team_id=team.id, user_id=user.id, role="leader", status="active")
    )
    kb = KnowledgeBase(team_id=team.id, name="Parsing KB", created_by=user.id)
    session.add(kb)
    session.flush()

    path = tmp_path / file_name
    document = Document(
        knowledge_base_id=kb.id,
        owner_user_id=user.id,
        file_name=file_name,
        source_uri=str(path),
    )
    session.add(document)
    session.flush()
    version = DocumentVersion(
        document_id=document.id,
        content_fingerprint="fingerprint",
        file_size_bytes=0,
    )
    session.add(version)
    session.flush()
    document.current_version_id = version.id
    run = IngestionRun(document_version_id=version.id, status="queued")
    session.add(run)
    session.commit()
    return session, document, run, settings, path


def test_processor_ingests_docx_content(tmp_path, monkeypatch):
    session, document, run, settings, path = _ingestion_context(tmp_path, "guide.docx")
    docx = DocxDocument()
    docx.add_heading("Operations", level=1)
    docx.add_paragraph("Restart the worker before rebuilding the index.")
    docx.save(path)
    monkeypatch.setattr(
        processor_module.EmbeddingClient,
        "embed",
        lambda self, texts: [[0.1, 0.2] for _ in texts],
    )

    chunk_ids = process_document(document, run, session, settings)

    chunks = session.query(ChunkRecord).all()
    assert chunk_ids
    assert run.status == "completed"
    assert any("Restart the worker" in chunk.content for chunk in chunks)
    session.close()


def test_processor_persists_stable_source_locator_fields(tmp_path, monkeypatch):
    session, document, run, settings, path = _ingestion_context(tmp_path, "guide.md")
    path.write_text("# Operations\n\nRestart the worker before rebuilding.", encoding="utf-8")
    monkeypatch.setattr(
        processor_module.EmbeddingClient,
        "embed",
        lambda self, texts: [[0.1, 0.2] for _ in texts],
    )

    chunk_ids = process_document(document, run, session, settings)

    assert chunk_ids
    chunk = session.query(ChunkRecord).one()
    locator = json.loads(chunk.source_locator)
    knowledge_base = session.get(KnowledgeBase, document.knowledge_base_id)
    assert knowledge_base is not None
    download_url = urlparse(locator["source_url"])
    assert download_url.path == "/api/v1/docs/download"
    assert parse_qs(download_url.query) == {
        "team_id": [knowledge_base.team_id],
        "document_id": [document.id],
    }
    assert locator["document_id"] == document.id
    assert locator["document_version_id"] == document.current_version_id
    assert locator["file_name"] == "guide.md"
    assert locator["source_path"] == "guide.md"
    assert locator["heading_path"] == ["Operations"]
    session.close()


def test_processor_marks_unreadable_document_as_failed(tmp_path, monkeypatch):
    session, document, run, settings, path = _ingestion_context(tmp_path, "broken.pdf")
    path.write_bytes(b"not a valid pdf")
    monkeypatch.setattr(
        processor_module.EmbeddingClient,
        "embed",
        lambda self, texts: [[0.1, 0.2] for _ in texts],
    )

    chunk_ids = process_document(document, run, session, settings)

    assert chunk_ids == []
    assert run.status == "failed"
    assert run.error_code == "DOCUMENT_PARSE_FAILED"
    assert document.doc_status == "failed"
    session.close()


def test_processor_does_not_overwrite_cancelled_run(tmp_path):
    session, document, run, settings, path = _ingestion_context(tmp_path, "cancelled.md")
    path.write_text("# Cancelled\nShould not be processed.", encoding="utf-8")
    run.status = "cancelled"
    run.current_stage = "cancelled"
    session.commit()

    chunk_ids = process_document(document, run, session, settings)

    assert chunk_ids == []
    assert run.status == "cancelled"
    assert run.current_stage == "cancelled"
    assert document.doc_status == "building"
    assert session.query(ChunkRecord).count() == 0
    session.close()


def test_processor_honors_cancellation_that_arrives_during_embedding(
    tmp_path, monkeypatch
):
    session, document, run, settings, path = _ingestion_context(
        tmp_path, "cancel-during-run.md"
    )
    path.write_text("# Cancellation\nStop before indexing.", encoding="utf-8")
    engine = session.get_bind()
    session_factory = build_session_factory(engine)

    def cancel_then_embed(self, texts):
        with session_factory() as concurrent:
            concurrent_run = concurrent.get(IngestionRun, run.id)
            concurrent_run.status = "cancelled"
            concurrent_run.current_stage = "cancelled"
            concurrent_run.retryable = True
            concurrent.commit()
        return [[0.1, 0.2] for _ in texts]

    monkeypatch.setattr(
        processor_module.EmbeddingClient,
        "embed",
        cancel_then_embed,
    )

    chunk_ids = process_document(document, run, session, settings)

    assert chunk_ids == []
    assert run.status == "cancelled"
    assert run.current_stage == "cancelled"
    assert session.query(ChunkRecord).count() == 0
    session.close()
